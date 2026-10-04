"""Conversion and editing tools (no Flask code, so they are easy to test).

Every tool returns (data, ext, info): the file's bytes, its extension (".pdf", ".zip", ...)
and a small dict {"title": ..., "lines": [...]} that the preview page displays.
"""
import math
import re
import zipfile
from io import BytesIO

import pypdfium2 as pdfium
from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas

from pdf_ops import (PdfEditError, format_pages, parse_pages, plural, read_pdf,
                     to_bytes)
from thumbs import LOCK as PDFIUM_LOCK


# ---------- helpers ----------
def human_size(n):
    if n < 1024:
        return f"{n} B"
    if n < 1024 ** 2:
        return f"{n / 1024:.0f} KB"
    return f"{n / 1024 ** 2:.1f} MB"


def make_zip(files):
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            z.writestr(name, data)
    return buf.getvalue()


def pdf_bytes(writer):
    return to_bytes(writer).getvalue()


def pdf_from_pages(reader, pages):
    writer = PdfWriter()
    for n in pages:
        writer.add_page(reader.pages[n - 1])
    return pdf_bytes(writer)


def render_page(data, index, dpi, max_pixels=80_000_000):
    """Render page `index` (0-based) to an RGB PIL image."""
    with PDFIUM_LOCK:
        pdf = pdfium.PdfDocument(data)
        try:
            page = pdf[index]
            w, h = page.get_size()
            scale = min(dpi / 72, math.sqrt(max_pixels / max(w * h, 1)))
            return page.render(scale=scale).to_pil().convert("RGB")
        finally:
            pdf.close()


def clip_lines(lines, limit=15):
    if len(lines) > limit:
        return lines[:limit] + [f"… and {len(lines) - limit} more"]
    return lines


# ---------- split / extract ----------
def parse_groups(spec, total):
    """'1-3, 5, 7-9' -> [[1, 2, 3], [5], [7, 8, 9]] (one group per comma-separated part)."""
    parts = [p for p in re.split(r"[,;]+", spec or "") if p.strip()]
    if not parts:
        raise PdfEditError("Please enter at least one page or range, e.g. 1-3, 5, 7-9.")
    return [parse_pages(p, total) for p in parts]


def split(reader, stem, mode, ranges="", pages="", every=1):
    total = len(reader.pages)
    if mode == "extract":
        chosen = parse_pages(pages, total)
        info = {"title": f"Extracted {plural(len(chosen))}",
                "lines": [f"Pages: {format_pages(chosen)}", f"Total: {plural(len(chosen))}"]}
        return pdf_from_pages(reader, chosen), ".pdf", info

    if mode == "each":
        groups = [[i] for i in range(1, total + 1)]
    elif mode == "every":
        if not 1 <= every < total:
            raise PdfEditError(f"Enter a number between 1 and {total - 1}.")
        groups = [list(range(i, min(i + every, total + 1))) for i in range(1, total + 1, every)]
    elif mode == "ranges":
        groups = parse_groups(ranges, total)
    else:
        raise PdfEditError("Unknown split option.")

    width = len(str(len(groups)))
    files, lines = [], []
    for i, group in enumerate(groups, start=1):
        label = f"page_{group[0]}" if len(group) == 1 else f"pages_{group[0]}-{group[-1]}"
        name = f"{stem}_{i:0{width}d}_{label}.pdf"
        files.append((name, pdf_from_pages(reader, group)))
        lines.append(f"{name} — {plural(len(group))}")
    info = {"title": f"Split into {plural(len(files), 'file')}", "lines": clip_lines(lines)}
    if len(files) == 1:
        return files[0][1], ".pdf", info
    return make_zip(files), ".zip", info


# ---------- rotate ----------
def rotate(reader, rotations):
    """rotations: {page_number: degrees clockwise (multiple of 90)}."""
    total = len(reader.pages)
    writer = PdfWriter()
    changed = []
    for i, page in enumerate(reader.pages, start=1):
        degrees = rotations.get(i, 0)
        if degrees % 360:
            page.rotate(degrees % 360)
            changed.append(i)
        writer.add_page(page)
    if not changed:
        raise PdfEditError("No pages were rotated.")
    info = {"title": f"Rotated {plural(len(changed))}",
            "lines": [f"Pages: {format_pages(changed)}", f"Total: {plural(total)}"]}
    return pdf_bytes(writer), ".pdf", info


# ---------- compress ----------
COMPRESS_LEVELS = {            # level: (JPEG quality, longest image side in pixels)
    "low": (85, 3500),
    "recommended": (70, 2200),
    "extreme": (45, 1500),
}


def compress(data, level, name="file"):
    quality, max_side = COMPRESS_LEVELS[level]
    reader = read_pdf(BytesIO(data), name)
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)

    recompressed = 0
    for page in writer.pages:
        try:
            page.compress_content_streams()
        except Exception:
            pass
        try:
            images = list(page.images)
        except Exception:
            images = []
        for image in images:
            try:
                pil = image.image
                if pil.mode not in ("RGB", "L"):          # skip masks, CMYK, transparency...
                    continue
                if max(pil.size) > max_side:
                    pil.thumbnail((max_side, max_side))
                image.replace(pil, quality=quality)
                recompressed += 1
            except Exception:
                continue
    try:
        writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
    except Exception:
        pass

    new = pdf_bytes(writer)
    old_size, new_size = len(data), len(new)
    if new_size >= old_size:
        info = {"title": "This PDF is already well optimized",
                "lines": [f"Size: {human_size(old_size)}",
                          "Compressing it further would not make it smaller, so the original is kept."]}
        return data, ".pdf", info
    saved = 100 * (old_size - new_size) / old_size
    info = {"title": f"Made {saved:.0f}% smaller",
            "lines": [f"Before: {human_size(old_size)}", f"After: {human_size(new_size)}",
                      f"Images recompressed: {recompressed}"]}
    return new, ".pdf", info


# ---------- protect / unlock ----------
def protect(reader, password):
    if not password:
        raise PdfEditError("Please enter a password.")
    writer = PdfWriter(clone_from=reader)
    writer.encrypt(user_password=password, owner_password=password, algorithm="AES-256")
    info = {"title": "PDF protected with a password",
            "lines": ["Encryption: AES-256",
                      "Anyone opening the file will need the password — keep it safe, it can't be recovered."]}
    return pdf_bytes(writer), ".pdf", info


def unlock(data, password):
    try:
        reader = PdfReader(BytesIO(data))
    except Exception:
        raise PdfEditError("This is not a valid PDF file.")
    if not reader.is_encrypted:
        raise PdfEditError("This PDF is not password protected.")
    if not reader.decrypt(password or ""):
        raise PdfEditError("That password is not correct.")
    writer = PdfWriter(clone_from=reader)
    info = {"title": "Password removed",
            "lines": [f"Pages: {len(writer.pages)}", "The new file opens without a password."]}
    return pdf_bytes(writer), ".pdf", info


# ---------- overlays: page numbers + watermark ----------
def _prepare_pages(reader):
    """Copy pages into a writer; bake any /Rotate into the content so overlays line up."""
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    for page in writer.pages:
        if page.rotation:
            page.transfer_rotation_to_content()
    return writer


def _apply_overlay(writer, draw):
    pages = list(writer.pages)
    buf = BytesIO()
    c = canvas.Canvas(buf)
    for i, page in enumerate(pages):
        x0, y0 = float(page.mediabox.left), float(page.mediabox.bottom)
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        c.setPageSize((w, h))
        c.saveState()
        c.translate(x0, y0)
        draw(c, i, w, h)
        c.restoreState()
        c.showPage()
    c.save()
    buf.seek(0)
    overlay = PdfReader(buf)
    for page, over in zip(pages, overlay.pages):
        page.merge_page(over)


PAGE_NUMBER_POSITIONS = {
    "bottom-center": ("bottom", "center"), "bottom-right": ("bottom", "right"),
    "bottom-left": ("bottom", "left"), "top-center": ("top", "center"),
    "top-right": ("top", "right"), "top-left": ("top", "left"),
}


def add_page_numbers(reader, position, fmt, start, skip_first):
    total = len(reader.pages)
    if skip_first and total < 2:
        raise PdfEditError("The PDF has only one page, so there is nothing to number after skipping it.")
    vertical, align = PAGE_NUMBER_POSITIONS[position]
    numbered = total - (1 if skip_first else 0)
    last_number = start + numbered - 1
    writer = _prepare_pages(reader)

    def draw(c, i, w, h):
        if skip_first and i == 0:
            return
        n = start + i - (1 if skip_first else 0)
        text = {"number": f"{n}", "page": f"Page {n}",
                "page_of": f"Page {n} of {last_number}"}[fmt]
        size, margin = 11, 28
        c.setFont("Helvetica", size)
        c.setFillColorRGB(0.1, 0.1, 0.1)
        y = margin if vertical == "bottom" else h - margin - size * 0.8
        if align == "left":
            c.drawString(margin, y, text)
        elif align == "right":
            c.drawRightString(w - margin, y, text)
        else:
            c.drawCentredString(w / 2, y, text)

    _apply_overlay(writer, draw)
    info = {"title": f"Numbered {plural(numbered)}",
            "lines": [f"Position: {position.replace('-', ' ')}",
                      f"Numbers: {start} to {last_number}"]}
    return pdf_bytes(writer), ".pdf", info


WATERMARK_COLORS = {"gray": (0.45, 0.45, 0.45), "red": (0.85, 0.1, 0.1),
                    "blue": (0.1, 0.3, 0.85), "black": (0, 0, 0)}
WATERMARK_OPACITY = {"light": 0.15, "medium": 0.3, "strong": 0.5}
WATERMARK_SIZE = {"small": 0.45, "medium": 0.7, "large": 1.0}
FONT = "Helvetica-Bold"


def add_watermark(reader, text, size, opacity, color, angle):
    text = (text or "").strip()
    if not text:
        raise PdfEditError("Please enter the watermark text.")
    if len(text) > 60:
        raise PdfEditError("The watermark text is too long (60 characters at most).")
    try:
        text.encode("latin-1")
    except UnicodeEncodeError:
        raise PdfEditError("Watermarks currently support Latin letters, numbers and common symbols only.")

    rgb, alpha = WATERMARK_COLORS[color], WATERMARK_OPACITY[opacity]
    width_at_100 = pdfmetrics.stringWidth(text, FONT, 100)
    writer = _prepare_pages(reader)

    def draw(c, i, w, h):
        length = (1.1 * min(w, h) if angle else 0.8 * w) * WATERMARK_SIZE[size]
        font_size = length / width_at_100 * 100
        c.setFillColorRGB(*rgb)
        c.setFillAlpha(alpha)
        c.setFont(FONT, font_size)
        c.translate(w / 2, h / 2)
        c.rotate(angle)
        c.drawCentredString(0, -font_size / 3, text)

    _apply_overlay(writer, draw)
    info = {"title": "Watermark added",
            "lines": [f"Text: {text}", f"Applied to {plural(len(writer.pages))}"]}
    return pdf_bytes(writer), ".pdf", info


# ---------- images -> PDF ----------
PAGE_SIZES = {"a4": (595, 842), "letter": (612, 792)}
MARGINS = {"none": 0, "small": 18, "big": 36}


def _to_rgb(im):
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        im = im.convert("RGBA")
        background = Image.new("RGBA", im.size, "white")
        im = Image.alpha_composite(background, im)
    return im.convert("RGB")


def images_to_pdf(images, page_size, orientation, margin):
    """images: [(name, bytes)]. page_size: fit|a4|letter. orientation: auto|portrait|landscape."""
    if not images:
        raise PdfEditError("Please add at least one image.")
    writer = PdfWriter()
    margin_pt = MARGINS[margin]
    for name, data in images:
        try:
            with Image.open(BytesIO(data)) as im:
                im = _to_rgb(ImageOps.exif_transpose(im))
        except Exception:
            raise PdfEditError(f"'{name}' is not a supported image.")
        iw, ih = im.size

        if page_size == "fit":
            dpi = max(150, round(max(iw, ih) * 72 / 1000))     # keep pages a sensible size
            pad = round(margin_pt * dpi / 72)
            page = Image.new("RGB", (iw + 2 * pad, ih + 2 * pad), "white")
            page.paste(im, (pad, pad))
        else:
            dpi = 200
            pw, ph = PAGE_SIZES[page_size]
            landscape = (iw > ih) if orientation == "auto" else orientation == "landscape"
            if landscape:
                pw, ph = ph, pw
            px = lambda pt: round(pt * dpi / 72)
            page = Image.new("RGB", (px(pw), px(ph)), "white")
            box_w, box_h = px(pw - 2 * margin_pt), px(ph - 2 * margin_pt)
            scale = min(box_w / iw, box_h / ih)
            fitted = im.resize((max(1, round(iw * scale)), max(1, round(ih * scale))), Image.LANCZOS)
            page.paste(fitted, ((px(pw) - fitted.width) // 2, (px(ph) - fitted.height) // 2))

        buf = BytesIO()
        page.save(buf, "PDF", resolution=dpi, quality=90)
        buf.seek(0)
        writer.add_page(PdfReader(buf).pages[0])

    lines = [f"{i}. {name}" for i, (name, _) in enumerate(images, start=1)]
    info = {"title": f"Created a PDF from {plural(len(images), 'image')}", "lines": clip_lines(lines)}
    return pdf_bytes(writer), ".pdf", info


# ---------- PDF -> images ----------
def pdf_to_images(data, stem, total, fmt, dpi, pages_spec=""):
    pages = sorted(set(parse_pages(pages_spec, total))) if pages_spec.strip() else list(range(1, total + 1))
    width = len(str(total))
    files = []
    for n in pages:
        image = render_page(data, n - 1, dpi)
        buf = BytesIO()
        if fmt == "jpg":
            image.save(buf, "JPEG", quality=90, dpi=(dpi, dpi))
            ext = ".jpg"
        else:
            image.save(buf, "PNG", dpi=(dpi, dpi))
            ext = ".png"
        files.append((f"{stem}_page_{n:0{width}d}{ext}", buf.getvalue()))
    info = {"title": f"Converted {plural(len(files))} to {fmt.upper()}",
            "lines": [f"Resolution: {dpi} DPI", f"Pages: {format_pages(pages)}"]}
    if len(files) == 1:
        return files[0][1], files[0][0][-4:], info
    return make_zip(files), ".zip", info


# ---------- PDF -> Word / text ----------
def extract_text_pages(data, total):
    pages = []
    for i in range(total):
        with PDFIUM_LOCK:
            pdf = pdfium.PdfDocument(data)
            try:
                text = pdf[i].get_textpage().get_text_range()
            finally:
                pdf.close()
        pages.append(text.replace("\r\n", "\n").replace("\r", "\n"))
    return pages


def _paragraphs(text):
    """Join hard-wrapped lines back into paragraphs (simple heuristic)."""
    lines = [ln.strip() for ln in text.split("\n")]
    longest = max((len(ln) for ln in lines), default=0)
    paragraphs, current, previous = [], [], ""
    for ln in lines:
        if not ln:
            if current:
                paragraphs.append(" ".join(current))
                current = []
            previous = ""
            continue
        if current and (len(previous) < 0.6 * longest or previous.endswith((".", ":", "!", "?"))):
            paragraphs.append(" ".join(current))
            current = []
        current.append(ln)
        previous = ln
    if current:
        paragraphs.append(" ".join(current))
    return paragraphs


_BAD_CHARS = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def _clean(text):
    return _BAD_CHARS.sub("", text.encode("utf-8", "ignore").decode("utf-8"))


def pdf_to_word(data, total, fmt):
    pages = [_clean(t) for t in extract_text_pages(data, total)]
    if sum(len(t.strip()) for t in pages) < 5:
        raise PdfEditError("This PDF has no selectable text — it looks like a scan. "
                           "Run it through the OCR PDF tool first, then convert it.")
    paragraphs_per_page = [_paragraphs(t) for t in pages]
    words = sum(len(p.split()) for page in paragraphs_per_page for p in page)

    if fmt == "txt":
        body = "\n\n\n".join("\n\n".join(page) for page in paragraphs_per_page)
        out, ext = body.encode("utf-8"), ".txt"
    else:
        from docx import Document
        doc = Document()
        for i, page in enumerate(paragraphs_per_page):
            if i:
                doc.add_page_break()
            for paragraph in page:
                doc.add_paragraph(paragraph)
        buf = BytesIO()
        doc.save(buf)
        out, ext = buf.getvalue(), ".docx"
    info = {"title": f"Converted {plural(total)} ({words:,} words)",
            "lines": ["Text is extracted and re-flowed into paragraphs.",
                      "Complex layouts, tables and images are not kept."]}
    return out, ext, info
